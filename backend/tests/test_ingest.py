import io

import docx
import pytest

from coach.ingest import Document, IngestError, extract_text, fit_to_budget


def make_pdf(text: str) -> bytes:
    """Smallest valid one-page PDF with a text run, so tests need no PDF library."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def test_text_strips_bom_and_collapses_whitespace():
    raw = "\ufeffJordan   Rivera\r\n\r\n\r\n\r\nPython\t\tSQL".encode("utf-8")
    assert extract_text("resume.txt", raw) == "Jordan Rivera\n\nPython SQL"


def test_markdown_is_read_as_text():
    assert extract_text("notes.MD", b"# Skills\n- Python") == "# Skills\n- Python"


def test_docx_paragraphs_and_tables():
    document = docx.Document()
    document.add_paragraph("Software Engineering Intern")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "4 years"
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text("resume.docx", buffer.getvalue())
    assert "Software Engineering Intern" in text
    assert "Python | 4 years" in text


def test_pdf_text_is_extracted():
    assert "Built an ETL job" in extract_text("resume.pdf", make_pdf("Built an ETL job"))


@pytest.mark.parametrize("name", ["photo.png", "resume", "data.csv"])
def test_unsupported_types_are_rejected(name):
    with pytest.raises(IngestError, match="unsupported"):
        extract_text(name, b"hello")


def test_empty_files_are_rejected():
    with pytest.raises(IngestError, match="no readable text"):
        extract_text("blank.txt", b"   \n  ")


def test_corrupt_pdf_is_a_clean_error():
    with pytest.raises(IngestError, match="could not be read"):
        extract_text("broken.pdf", b"not really a pdf")


def test_budget_keeps_short_docs_whole_and_trims_long_ones():
    docs = [
        Document("resume.md", "file", "1" * 5000),
        Document("Achievements", "text", "2" * 300),
        Document("portfolio.md", "file", "3" * 9000),
    ]
    corpus, sources = fit_to_budget(docs, 4000)
    by_name = {s.name: s for s in sources}
    assert not by_name["Achievements"].truncated
    assert by_name["resume.md"].truncated and by_name["portfolio.md"].truncated
    assert corpus.count("2") == 300
    assert corpus.count("1") + corpus.count("2") + corpus.count("3") == 4000
    assert corpus.count("1") == corpus.count("3") == 1850  # leftover budget split evenly


def test_budget_untouched_when_everything_fits():
    corpus, sources = fit_to_budget([Document("a.txt", "file", "hello world")], 100)
    assert corpus == "### a.txt\nhello world"
    assert sources[0].chars == 11 and not sources[0].truncated
