"""Turn uploaded files and pasted text into plain text the model can read."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import PurePath

from .schemas import SourceDoc

ALLOWED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")

_SPACES = re.compile(r"[ \t\f\v]+")
_BLANK_LINES = re.compile(r"\n{3,}")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0e-\x1f\x7f]")


class IngestError(ValueError):
    """One uploaded file could not be read."""


@dataclass
class Document:
    name: str
    kind: str  # "file" or "text"
    text: str


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").lstrip("\ufeff")
    text = _CONTROL.sub(" ", text)
    text = "\n".join(_SPACES.sub(" ", line).strip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()


def extract_text(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise IngestError(f"{filename}: unsupported file type. Use PDF, DOCX, TXT or MD.")
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            raw = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            raw = "\n".join(parts)
        else:
            raw = data.decode("utf-8-sig", errors="replace")
    except Exception as exc:  # corrupt, encrypted or mislabeled files
        raise IngestError(f"{filename}: could not be read ({exc.__class__.__name__}).") from exc
    text = normalize(raw)
    if not text:
        raise IngestError(f"{filename}: no readable text. Scanned PDFs need OCR; paste the text instead.")
    return text


def fit_to_budget(docs: list[Document], budget: int) -> tuple[str, list[SourceDoc]]:
    """Join documents into one prompt section that fits a character budget.

    Short documents stay whole and the leftover budget is split evenly across
    the long ones, so one huge file cannot crowd out everything else.
    """
    order = sorted(range(len(docs)), key=lambda i: len(docs[i].text))
    allowance: dict[int, int] = {}
    remaining = budget
    for position, index in enumerate(order):
        share = remaining // (len(order) - position)
        allowance[index] = min(len(docs[index].text), share)
        remaining -= allowance[index]

    sections: list[str] = []
    sources: list[SourceDoc] = []
    for index, doc in enumerate(docs):
        keep = allowance[index]
        truncated = keep < len(doc.text)
        body = doc.text[:keep].rstrip() + (" [...truncated]" if truncated else "")
        sections.append(f"### {doc.name}\n{body}")
        sources.append(SourceDoc(name=doc.name, kind=doc.kind, chars=len(doc.text), truncated=truncated))
    return "\n\n".join(sections), sources
